import json, pandas as pd
from screen_lpe import SCREEN
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FUNC=set('및 하다 되다 있다 위하다 통하다 대하다 따르다 관하다 의하다 인하다 같다 이상 이하 이후 이전 등'.split())
cands={c['form']:c for c in json.load(open('inductive.json',encoding='utf-8'))['cands'] if c['form'] not in FUNC}
NV=json.load(open('inductive.json'))['n_valid']
ded=pd.read_excel('McLeod_TableII_한국어후보사전_초안.xlsx','후보사전_원어별')
eng={}
for _,r in ded.iterrows():
    eng.setdefault(str(r['한국어 후보']),[]).append(f"{r['원어(Table II)']}")
rows=[]; missing=[]; seen=set()
for dim,themes in SCREEN.items():
    for th,ws in themes.items():
        for w in ws.split():
            flag=w.startswith('!'); w=w.lstrip('!')
            if (dim,w) in seen: continue
            seen.add((dim,w))
            c=cands.get(w)
            if not c: missing.append(w); continue
            rows.append(dict(dim=dim,theme=th,w=w,flag=flag,c=c,eng=', '.join(sorted(set(eng.get(w,[]))))))
order={'Logos':0,'Pathos':1,'Ethos':2}
rows.sort(key=lambda r:(order[r['dim']],-r['c']['docs3']))
SECKR={"ITEM":"아이템개요","MOTIV":"개발동기","DOM":"국내시장","FOR":"해외시장"}

wb=load_workbook('DIPS_귀납후보_단어목록.xlsx')
old=wb['귀납후보']; old.title='전체빈도목록(참고)'
ws=wb.create_sheet('LPE후보',1)
A='Arial'; hf=Font(name=A,bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='1F3864')
thin=Side(style='thin',color='BFBFBF'); B=Border(left=thin,right=thin,top=thin,bottom=thin); base=Font(name=A,size=10)
fills={'Logos':'DDEBF7','Pathos':'FCE4D6','Ethos':'E2EFDA'}
cols=[('No',6),('제안 차원',10),('테마',16),('단어(원형)',12),('Table II 원어',18),('문맥 확인',9),('3회↑ 문서 수',9),('문서 비율',9),('총빈도',8),('예문 항목',9),('예문',70),('평정자1',10),('평정자2',10),('평정자3',10),('판정',10),('비고',18)]
for j,(h,w) in enumerate(cols,1):
    c=ws.cell(row=1,column=j,value=h); c.font=hf; c.fill=hfill; c.border=B; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.column_dimensions[get_column_letter(j)].width=w
dv=DataValidation(type='list',formula1='"Logos,Pathos,Ethos,제외"',allow_blank=True); ws.add_data_validation(dv)
inp=PatternFill('solid',fgColor='FFFF00')
for i,r in enumerate(rows,2):
    c=r['c']; kw=c['kwic']
    vals=[i-1,r['dim'],r['theme'],r['w'],r['eng'],'필요' if r['flag'] else '',c['docs3'],f'=G{i}/{NV}',c['total'],SECKR[kw[1]] if kw else '',kw[2] if kw else '']
    for j,v in enumerate(vals,1):
        x=ws.cell(row=i,column=j,value=v); x.font=base; x.border=B; x.alignment=Alignment(vertical='top',wrap_text=(j==11))
    ws.cell(row=i,column=8).number_format='0.0%'
    ws.cell(row=i,column=2).fill=PatternFill('solid',fgColor=fills[r['dim']])
    if r['flag']: ws.cell(row=i,column=6).fill=PatternFill('solid',fgColor='F8CBAD')
    for j in (12,13,14):
        x=ws.cell(row=i,column=j); x.fill=inp; x.border=B; x.font=base
    f=ws.cell(row=i,column=15,value=f'=IF(COUNTA(L{i}:N{i})<3,"",IF(AND(L{i}=M{i},M{i}=N{i}),L{i},"불일치"))'); f.font=base; f.border=B
    ws.cell(row=i,column=16).border=B
dv.add(f'L2:N{len(rows)+1}'); ws.freeze_panes='E2'; ws.auto_filter.ref=f'A1:P{len(rows)+1}'

# README 갱신
rd=wb['README']; rd.delete_rows(1,rd.max_row)
lines=[("DIPS 사업계획서 귀납 후보: Logos·Pathos·Ethos 선별본",True),("",False),
(f"LPE후보 시트: 전체 빈도목록 {len(cands):,}개에서 25개 테마(Table II)에 해당할 수 있는 단어만 고른 {len(rows)}개. 원 논문에서 귀납 단계 출력 중 'logos·pathos·ethos와 관련 있다고 본 단어만 추가'한 단계에 해당.",False),
("1차 선별은 Claude가 한 것이며 확정이 아님. 평정자 3인이 노란 칸에 Logos/Pathos/Ethos/제외를 입력하고, 세 명이 같은 차원을 고른 단어만 남김.",False),
("제안 차원·테마: 원 논문 Table II의 테마 기준. Table II 원어: 같은 단어가 연역 후보(Table II 번역)에도 있으면 대응 원어.",False),
("문맥 확인 '필요': 사업계획서에서 주로 기술·제조 용어로 쓰이는 단어(공정=工程, 품질, 인식, 기술, 시스템 등). 예문을 보고 수사로 쓰였는지 판단.",False),
("",False),
("전체빈도목록(참고) 시트: 문서마다 3회 이상 쓰인 모든 형태소. 대부분 로봇·업종 주제어(하중, 타임, 저장 등)라 후보가 아님. 감사·재현용으로만 남김.",False),
("정제 규칙: PDF 줄바꿈으로 단어가 끊긴 조각(다\\n양하게 → '다')을 이어 붙인 뒤 분석, 1음절 명사·부사 제외, 기능어(및·하다·위하다 등) 제외.",False),
("다른 시트: 문서별_추출현황 / 항목텍스트.",False)]
for i,(t,b) in enumerate(lines,1):
    x=rd.cell(row=i,column=1,value=t); x.font=Font(name=A,bold=b,size=12 if i==1 else 10); x.alignment=Alignment(wrap_text=True,vertical='top')
wb.save('DIPS_LPE후보_단어목록.xlsx')
from collections import Counter
print(len(rows), Counter(r['dim'] for r in rows), 'missing:', missing)
