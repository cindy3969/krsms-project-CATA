import json, re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ind = json.load(open("/home/claude/inductive.json", encoding="utf-8"))
secs = json.load(open("/home/claude/sections.json", encoding="utf-8"))
ded = pd.read_excel("/home/claude/McLeod_TableII_한국어후보사전_초안.xlsx", sheet_name="고유키_요약")

# 연역 후보(단일 형태소 키)와 매칭
dmap = {}
for _, r in ded.iterrows():
    key = str(r["매칭 키"])
    if " + " in key or "/" not in key:
        continue
    form, tag = key.split("/")
    if tag in ("VV", "VA"):
        form = form + "다"
    dmap.setdefault((form, tag), set()).add(r["차원"].split("(")[0])

FUNC = set("및 하다 되다 있다 위하다 통하다 대하다 따르다 관하다 의하다 인하다 같다 이상 이하 이후 이전 등".split())
cands = [c for c in ind["cands"] if c["form"] not in FUNC]
cands.sort(key=lambda c: (-c["docs3"], -c["total"]))
NV = ind["n_valid"]

ARIAL = "Arial"
hdr_font = Font(name=ARIAL, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F3864")
inp = PatternFill("solid", fgColor="FFFF00")
hit = PatternFill("solid", fgColor="E2EFDA")
tmpl = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
base = Font(name=ARIAL, size=10)

def header(ws, cols):
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border = hdr_font, hdr_fill, B
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32

wb = Workbook()
ws = wb.active; ws.title = "README"
ok_docs = sum(1 for d in secs if sum(len(d[s]) >= 150 for s in ["ITEM", "MOTIV", "DOM", "FOR"]) >= 3)
lines = [
 ("DIPS 사업계획서 귀납 후보 단어 목록 (McLeod et al., 2022 절차 ③단계)", True),
 ("", False),
 (f"표본: 사업계획서 {len(secs)}건 (2025 신규지원 {sum(d['year']=='2025' for d in secs)}건, 2026 신규지원 {sum(d['year']=='2026' for d in secs)}건). 4개 항목 중 3개 이상 추출된 {ok_docs}건으로 후보를 셈.", False),
 ("분석 항목: ① 아이템(기술) 개요(개발 동기 제외: 소개·차별성·BM·개선과제) ② 창업아이템 개발 동기 및 목적 ③ 국내(내수) 시장 진출 현황 및 계획 ④ 해외시장 진출 현황 및 계획. 양식 안내문(※), 표 머리말, 쪽번호는 제거.", False),
 ("규칙: 원 논문처럼 '한 문서 안에서 3회 이상 쓰인 단어'를 문서마다 모아 합집합을 후보로 함. 단어 단위 = Kiwi 형태소 원형(명사 NNG, 동사·형용사 원형, 어근 XR, 부사 MAG). 고유명사(NNP)·영문·숫자는 제외. '및·하다·위하다' 같은 기능어는 제거.", False),
 (f"결과: 후보 {len(cands):,}개 (원 논문의 연역+귀납 후보 4,679개와 비슷한 규모). 이 중 연역 후보(Table II 한국어판)와 겹치는 단어는 초록색.", False),
 ("", False),
 ("열 설명", True),
 ("3회↑ 문서 수: 그 단어를 3번 이상 쓴 사업계획서 수. 문서 비율 = 이 값 ÷ 분석 문서 수. 등장 문서 수: 한 번이라도 쓴 문서 수. 총빈도: 전체 등장 횟수.", False),
 ("연역 매칭: 연역 후보 사전에서 같은 형태소 키가 배정된 차원. 비어 있으면 연역 단계에서는 없던 '새 단어'.", False),
 ("양식어휘: 시장·진출·현황처럼 양식 제목에 들어 있는 단어(회색). 모든 문서에 구조적으로 많이 나와서 수사 선택으로 보기 어려움 → 제외 검토 권장.", False),
 ("예문: 그 단어가 처음 나온 문장 하나(KWIC). 문맥을 보고 차원을 판단할 때 쓰기 위함.", False),
 ("", False),
 ("입력 방법 (노란 칸만 편집)", True),
 ("평정자1~3: 드롭다운에서 Logos / Pathos / Ethos / 제외 중 선택. 판정: 세 명이 같은 차원을 고르면 그 차원, 하나라도 다르면 '불일치', 미입력이면 빈칸.", False),
 ("권장 순서: '3회↑ 문서 수' 내림차순(이미 정렬됨)으로 위에서부터 평정. 원 논문은 전체를 평정했지만, 시간이 부족하면 문서 비율 1% 이상(3회↑ 문서 수 3 이상) 후보만 평정하고 그 기준을 보고.", False),
 ("입력 예시: 단어 '신뢰' → 평정자1 Ethos, 평정자2 Ethos, 평정자3 Ethos → 판정 Ethos", False),
 ("", False),
 ("다른 시트: 문서별_추출현황(항목별 글자 수·형태소 수, 확인 필요 문서 표시) / 항목텍스트(추출된 원문, 10% 잠재 내용 분석과 LLM 검증용).", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=ARIAL, bold=b, size=12 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

# ---- candidates ----
ws = wb.create_sheet("귀납후보")
cols = [("No", 6), ("형태소(원형)", 14), ("품사", 7), ("3회↑ 문서 수", 9), ("문서 비율", 9), ("등장 문서 수", 9), ("총빈도", 9),
        ("연역 매칭", 14), ("양식어휘", 8), ("예문 항목", 9), ("예문", 70), ("평정자1", 10), ("평정자2", 10), ("평정자3", 10), ("판정", 10), ("비고", 20)]
header(ws, cols)
dv = DataValidation(type="list", formula1='"Logos,Pathos,Ethos,제외"', allow_blank=True)
ws.add_data_validation(dv)
SECKR = {"ITEM": "아이템개요", "MOTIV": "개발동기", "DOM": "국내시장", "FOR": "해외시장"}
for i, c in enumerate(cands, 2):
    dims = "·".join(sorted(dmap.get((c["form"], c["tag"]), [])))
    kw = c["kwic"]
    vals = [i - 1, c["form"], c["tag"], c["docs3"], None, c["docs1"], c["total"], dims,
            "Y" if c["template"] else "", SECKR[kw[1]] if kw else "", kw[2] if kw else ""]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=j, value=v); cell.font = base; cell.border = B
        cell.alignment = Alignment(vertical="top", wrap_text=(j == 11))
    ws.cell(row=i, column=5, value=f"=D{i}/README_N").number_format = "0.0%"
    ws.cell(row=i, column=5).font = base; ws.cell(row=i, column=5).border = B
    if dims:
        for j in (2, 8): ws.cell(row=i, column=j).fill = hit
    if c["template"]:
        for j in (2, 9): ws.cell(row=i, column=j).fill = tmpl
    for j in (12, 13, 14):
        cell = ws.cell(row=i, column=j); cell.fill = inp; cell.border = B; cell.font = base
    f = ws.cell(row=i, column=15, value=f'=IF(COUNTA(L{i}:N{i})<3,"",IF(AND(L{i}=M{i},M{i}=N{i}),L{i},"불일치"))')
    f.font = base; f.border = B; f.alignment = Alignment(horizontal="center")
    ws.cell(row=i, column=16).border = B
last = len(cands) + 1
dv.add(f"L2:N{last}")
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:P{last}"

# named constant for N valid docs
from openpyxl.workbook.defined_name import DefinedName
wb.defined_names["README_N"] = DefinedName("README_N", attr_text="문서별_추출현황!$C$2")

# ---- per-doc ----
PDFCH = json.load(open("/home/claude/pdfchars.json", encoding="utf-8"))
def memo(d, nshort):
    if nshort <= 1:
        return "" if nshort == 0 else "1개 항목 미추출: 헤딩 확인"
    n, pages = PDFCH.get(d["file"], (None, None))
    if n == 0:
        return f"스캔본({pages}쪽, 텍스트 0자) → OCR 필요"
    if n is not None and n / max(pages, 1) < 400:
        return f"텍스트가 적음({n:,}자/{pages}쪽, 이미지 위주) → OCR 또는 수작업"
    return "헤딩 형식이 달라 항목 미추출 → 수작업 확인"
ws2 = wb.create_sheet("문서별_추출현황")
ws2["A1"] = "분석 문서 수(4개 항목 중 3개 이상 추출)"; ws2["A1"].font = Font(name=ARIAL, bold=True)
ws2["C2"] = f'=COUNTIF(N5:N{4+len(secs)},"OK")'; ws2["C2"].font = Font(name=ARIAL, bold=True)
ws2["A2"] = "분석 문서 수"; ws2["A2"].font = base
cols2 = [("연도", 7), ("트랙", 10), ("기업", 18), ("파일", 40), ("글자_아이템개요", 10), ("글자_개발동기", 10), ("글자_국내", 10), ("글자_해외", 10),
         ("형태소_아이템개요", 10), ("형태소_개발동기", 10), ("형태소_국내", 10), ("형태소_해외", 10), ("형태소_합계", 10), ("추출 상태", 10), ("메모", 30)]
for j, (h, w) in enumerate(cols2, 1):
    c = ws2.cell(row=4, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = B
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws2.column_dimensions[get_column_letter(j)].width = w
st = {r["file"]: r for r in ind["sec_tokens"]}
for i, d in enumerate(sorted(secs, key=lambda d: (d["year"], d["track"], d["firm"])), 5):
    t = st[d["file"]]
    nshort = sum(len(d[s]) < 150 for s in ["ITEM", "MOTIV", "DOM", "FOR"])
    vals = [d["year"], d["track"], d["firm"], d["file"], len(d["ITEM"]), len(d["MOTIV"]), len(d["DOM"]), len(d["FOR"]),
            t["tok_ITEM"], t["tok_MOTIV"], t["tok_DOM"], t["tok_FOR"], f"=SUM(I{i}:L{i})",
            "OK" if nshort <= 1 else "확인필요",
            memo(d, nshort)]
    for j, v in enumerate(vals, 1):
        c = ws2.cell(row=i, column=j, value=v); c.font = base; c.border = B
    if nshort:
        for j in (5, 6, 7, 8):
            if (j == 5 and len(d["ITEM"]) < 150) or (j == 6 and len(d["MOTIV"]) < 150) or (j == 7 and len(d["DOM"]) < 150) or (j == 8 and len(d["FOR"]) < 150):
                ws2.cell(row=i, column=j).fill = PatternFill("solid", fgColor="FCE4D6")
ws2.freeze_panes = "D5"
ws2.auto_filter.ref = f"A4:O{4+len(secs)}"

# ---- texts ----
ws3 = wb.create_sheet("항목텍스트")
cols3 = [("연도", 7), ("기업", 16), ("아이템(기술) 개요", 60), ("개발 동기 및 목적", 60), ("국내(내수) 시장", 60), ("해외시장", 60)]
header(ws3, cols3)
for i, d in enumerate(sorted(secs, key=lambda d: (d["year"], d["firm"])), 2):
    for j, v in enumerate([d["year"], d["firm"], d["ITEM"][:32000], d["MOTIV"][:32000], d["DOM"][:32000], d["FOR"][:32000]], 1):
        c = ws3.cell(row=i, column=j, value=v); c.font = base
        c.alignment = Alignment(vertical="top", wrap_text=(j > 2))
    ws3.row_dimensions[i].height = 120
ws3.freeze_panes = "C2"

out = "/home/claude/DIPS_귀납후보_단어목록.xlsx"
wb.save(out)
print(out, len(cands), "ded matches", sum(1 for c in cands if dmap.get((c['form'], c['tag']))))
