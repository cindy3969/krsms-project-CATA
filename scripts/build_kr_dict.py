import sys, re
sys.path.insert(0, "/home/claude")
from collections import defaultdict
from kr_seed_map import MAP, NOTES
from kiwipiepy import Kiwi
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

kiwi = Kiwi()
CONTENT = ("NNG", "NNP", "VV", "VA", "VX", "XR", "MAG", "XPN", "SL", "SN")
DIM_KR = {"logos": "Logos(논리)", "pathos": "Pathos(감정)", "ethos": "Ethos(성품)"}
OVERRIDE = {"신중": [("신중", "NNG")], "이점": [("이점", "NNG")], "증분": [("증분", "NNG")], "속상": [("속상", "XR")]}   # 단독 분석 오류 보정 (문장 속 분석 결과에 맞춤)

def analyze(term):
    # 명사형 후보는 문장 틀('그것은 X이다') 안에서 분석해야 실제 사업계획서 속 분석과 일치
    if term in OVERRIDE:
        toks = OVERRIDE[term]
    elif " " not in term and not term.endswith("다"):
        # 여러 문장 틀에서 분석해 단어를 가장 자연스럽게 복원하는 결과 선택
        best = None
        for pre, post in [("", "이다"), ("", "하다"), ("", "한 것이다"), ("", "을 보였다"), ("그것은 ", "이다")]:
            s = pre + term + post
            tt = [x for x in kiwi.tokenize(s) if len(pre) <= x.start < len(pre) + len(term)]
            forms = "".join(x.form for x in tt)
            bad = sum(x.tag.split("-")[0] in ("NNB", "NR", "MM", "NP", "EF", "JKS") for x in tt)
            sc = (forms == term, -bad, -len(tt))
            if best is None or sc > best[0]:
                best = (sc, tt)
        toks = [(x.form, x.tag.split("-")[0]) for x in best[1]]
    else:
        toks = [(x.form, x.tag.split("-")[0]) for x in kiwi.tokenize(term)]
    full = " + ".join(f"{f}/{g}" for f, g in toks)
    content = [f"{f}/{g}" for f, g in toks if g in CONTENT]
    return full, " + ".join(content), content

rows = []
for dim, themes in MAP.items():
    for theme, s in themes.items():
        for item in s.split(";"):
            item = item.strip()
            if not item: continue
            eng, kor = item.split("=", 1)
            for k in kor.split("/"):
                k = k.strip()
                full, key, content = analyze(k)
                rows.append(dict(dim=dim, theme=theme, eng=eng.strip(), kor=k, full=full, key=key, content=content))

# automatic flags
key_dims = defaultdict(set); key_engs = defaultdict(list)
for r in rows:
    key_dims[r["key"]].add(r["dim"]); key_engs[r["key"]].append(r["eng"])
for r in rows:
    f = []
    if len(r["content"]) == 1 and len(r["content"][0].split("/")[0]) == 1:
        f.append("1음절(과잉매칭)")
    joined = "".join(x.split("/")[0] for x in r["full"].split(" + ") if x.split("/")[-1] in CONTENT + ("XSN",))
    if " " not in r["kor"] and not r["kor"].endswith("다") and joined != r["kor"] and r["content"]:
        f.append("분석 확인 필요")
    if len(key_dims[r["key"]]) > 1:
        f.append("차원충돌:" + "·".join(sorted(key_dims[r["key"]])))
    elif len(set(key_engs[r["key"]])) > 1:
        f.append("차원내 중복")
    if len(r["content"]) > 1:
        f.append("구(句) 매칭")
    if len(r["content"]) == 0:
        f.append("분석실패")
    r["flags"] = ", ".join(f)
    r["note"] = NOTES.get(r["kor"], "")

wb = Workbook()
ARIAL = "Arial"
hdr_font = Font(name=ARIAL, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F3864")
input_fill = PatternFill("solid", fgColor="FFFF00")
flag_fill = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
base = Font(name=ARIAL, size=10)

# ---------------- README ----------------
ws = wb.active; ws.title = "README"
readme = [
 ("McLeod et al. (2022) Table II → 한국어 연역 후보 사전 (초안)", True),
 ("", False),
 ("출처: McLeod, Sears, Chandler, Payne & Brigham (2022), Journal of Management Studies 59(7), Table II (pp.1671–1673). 원문 891개 중 890개 전사(PDF 줄바꿈으로 1개 차이 가능).", False),
 ("성격: 원 논문 절차 중 ② 연역 단계(Rodale 유의어)에 해당하는 '후보' 목록. 한국어 번역은 Claude 초안이며 확정 사전이 아님.", False),
 ("다음 단계: ③ DIPS 사업계획서 귀납 후보 추가 → ④ 평정자 3인 판정(Holsti ≥ .70) → ⑤ 외부 전문가 추가 제안 + 5인 만장일치.", False),
 ("", False),
 ("시트 구성", True),
 ("후보사전_원어별: Table II 원어 1개당 한국어 후보 1행(후보가 여럿이면 여러 행). 평정자 판정 입력용.", False),
 ("고유키_요약: 형태소 매칭 키 기준으로 중복을 합친 목록. 실제 CATA 집계에 쓰일 사전의 모양.", False),
 ("집계: 차원별 후보 수·플래그 수·채택 수 (평정 입력에 따라 자동 갱신).", False),
 ("", False),
 ("열 설명", True),
 ("형태소 분석: Kiwi(kiwipiepy) 분석 결과 전체. 매칭 키: 내용 형태소(NNG·VV·VA·XR 등)만 남긴 것 = 사업계획서에서 이 키와 일치하는 형태소를 셈.", False),
 ("예: '증가예정' → 증가/NNG + 예정/NNG 로 분석되므로 키 '증가/NNG'가 잡힘. '도와' → 돕/VV 로 원형 복원.", False),
 ("접두사(XPN)는 키에 포함: '비호감' = 비/XPN + 호감/NNG 이므로 '호감'(ethos)과 구별됨. 집계 시 긴 키(구 매칭)를 먼저 세고, 그 형태소는 짧은 키로 중복 집계하지 않음.", False),
 ("매칭 키는 Kiwi가 실제 문장에서 내는 분석에 맞춘 것이라 언어학적으로 어색할 수 있음 (예: 무결성 → 무/XPN + 결성/NNG). 키를 고칠 땐 같은 단어를 넣은 예문을 Kiwi로 분석해 확인할 것.", False),
 ("자동 플래그: 1음절(과잉매칭) / 차원충돌(같은 키가 logos·pathos·ethos 중 둘 이상에 배정) / 차원내 중복 / 구(句) 매칭(형태소 2개 이상 연속).", False),
 ("검토 메모: 로봇·첨단제조 사업계획서 맥락에서 미리 짚어둔 위험(동음이의, 기술용어 오염 등).", False),
 ("", False),
 ("입력 방법 (노란 칸만 편집)", True),
 ("평정자1~3 열에 드롭다운으로 '유지' / '제외' / '차원변경' 중 선택. 차원변경이면 비고에 옮길 차원을 적음.", False),
 ("만장일치 열: 세 명 모두 '유지'면 '채택', 한 명이라도 다르면 '제외', 미입력이면 빈칸 (원 논문의 만장일치 규칙). 외부 전문가 2명을 더하려면 열을 추가해 같은 규칙으로 확장.", False),
 ("입력 예시: 평정자1=유지, 평정자2=유지, 평정자3=제외, 비고='제조 맥락에서 logos' → 만장일치=제외", False),
]
for i, (t, b) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=ARIAL, bold=b, size=12 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120

# ---------------- main sheet ----------------
ws = wb.create_sheet("후보사전_원어별")
headers = ["No", "차원", "테마(Table II)", "원어(Table II)", "한국어 후보", "형태소 분석(Kiwi)", "매칭 키",
           "자동 플래그", "검토 메모", "평정자1", "평정자2", "평정자3", "만장일치", "비고"]
widths = [6, 13, 14, 22, 16, 30, 22, 26, 48, 10, 10, 10, 11, 24]
for j, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(j)].width = w
dv = DataValidation(type="list", formula1='"유지,제외,차원변경"', allow_blank=True)
ws.add_data_validation(dv)
for i, r in enumerate(rows, 2):
    vals = [i - 1, DIM_KR[r["dim"]], r["theme"], r["eng"], r["kor"], r["full"], r["key"], r["flags"], r["note"]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = base; c.border = border
        c.alignment = Alignment(vertical="top", wrap_text=j in (6, 8, 9))
    if r["flags"] or r["note"]:
        for j in (8, 9): ws.cell(row=i, column=j).fill = flag_fill
    for j in (10, 11, 12):
        c = ws.cell(row=i, column=j); c.fill = input_fill; c.border = border; c.font = base
    c = ws.cell(row=i, column=13,
                value=f'=IF(COUNTA(J{i}:L{i})<3,"",IF(COUNTIF(J{i}:L{i},"유지")=3,"채택","제외"))')
    c.font = base; c.border = border; c.alignment = Alignment(horizontal="center")
    ws.cell(row=i, column=14).border = border
dv.add(f"J2:L{len(rows)+1}")
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:N{len(rows)+1}"
last = len(rows) + 1

# ---------------- unique key sheet ----------------
ws2 = wb.create_sheet("고유키_요약")
agg = {}
for r in rows:
    k = (r["dim"], r["key"])
    a = agg.setdefault(k, dict(engs=[], kors=set(), flags=set(), notes=set()))
    a["engs"].append(r["eng"]); a["kors"].add(r["kor"])
    if r["flags"]: a["flags"].update(r["flags"].split(", "))
    if r["note"]: a["notes"].add(r["note"])
h2 = ["차원", "매칭 키", "한국어 표기", "원어 수", "원어(Table II)", "자동 플래그", "검토 메모"]
w2 = [13, 24, 20, 8, 40, 30, 60]
for j, (h, w) in enumerate(zip(h2, w2), 1):
    c = ws2.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws2.column_dimensions[get_column_letter(j)].width = w
order = {"logos": 0, "pathos": 1, "ethos": 2}
for i, ((dim, key), a) in enumerate(sorted(agg.items(), key=lambda x: (order[x[0][0]], x[0][1])), 2):
    vals = [DIM_KR[dim], key, ", ".join(sorted(a["kors"])), len(a["engs"]), ", ".join(a["engs"]),
            ", ".join(sorted(a["flags"])), " / ".join(sorted(a["notes"]))]
    for j, v in enumerate(vals, 1):
        c = ws2.cell(row=i, column=j, value=v); c.font = base; c.border = border
        c.alignment = Alignment(vertical="top", wrap_text=j in (5, 6, 7))
    if a["flags"] or a["notes"]:
        for j in (6, 7): ws2.cell(row=i, column=j).fill = flag_fill
ws2.freeze_panes = "C2"
ws2.auto_filter.ref = f"A1:G{len(agg)+1}"
last2 = len(agg) + 1

# ---------------- summary sheet ----------------
ws3 = wb.create_sheet("집계")
h3 = ["차원", "Table II 원어 수", "한국어 후보 행 수", "고유 매칭 키 수", "플래그·메모 있는 행", "평정 완료 행", "채택(만장일치)"]
for j, h in enumerate(h3, 1):
    c = ws3.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws3.column_dimensions[get_column_letter(j)].width = 18
M = "후보사전_원어별"; U = "고유키_요약"
for i, dim in enumerate(["logos", "pathos", "ethos"], 2):
    d = DIM_KR[dim]
    n_eng = sum(len([x for x in s.split(";") if x.strip()]) for s in MAP[dim].values())
    vals = [d, n_eng,
            f"=COUNTIF({M}!$B$2:$B${last},A{i})",
            f"=COUNTIF({U}!$A$2:$A${last2},A{i})",
            f'=COUNTIFS({M}!$B$2:$B${last},A{i},{M}!$H$2:$H${last},"?*")+COUNTIFS({M}!$B$2:$B${last},A{i},{M}!$H$2:$H${last},"",{M}!$I$2:$I${last},"?*")',
            f'=COUNTIFS({M}!$B$2:$B${last},A{i},{M}!$M$2:$M${last},"채택")+COUNTIFS({M}!$B$2:$B${last},A{i},{M}!$M$2:$M${last},"제외")',
            f'=COUNTIFS({M}!$B$2:$B${last},A{i},{M}!$M$2:$M${last},"채택")']
    for j, v in enumerate(vals, 1):
        c = ws3.cell(row=i, column=j, value=v); c.font = base; c.border = border
        c.alignment = Alignment(horizontal="center")
ws3.cell(row=5, column=1, value="합계").font = Font(name=ARIAL, bold=True)
for j in range(2, 8):
    col = get_column_letter(j)
    c = ws3.cell(row=5, column=j, value=f"=SUM({col}2:{col}4)")
    c.font = Font(name=ARIAL, bold=True); c.border = border; c.alignment = Alignment(horizontal="center")
ws3.cell(row=7, column=1, value="Table II 원어 수는 원문에서 전사한 값(하드코딩; 원문 보고치 891).").font = Font(name=ARIAL, italic=True, size=9)

out = "/home/claude/McLeod_TableII_한국어후보사전_초안.xlsx"
wb.save(out)
print(out, "rows", len(rows), "unique keys", len(agg))
