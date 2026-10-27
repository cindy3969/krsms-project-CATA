"""연역 후보(표준국어대사전) -> 평정용 엑셀
열: 경로(씨앗/A 비슷한말/B 뜻풀이), 현대 사용 빈도(wordfreq zipf), 사업계획서 등장 문서 수
사용: python build_deductive_xlsx.py <expansion.json> <sections.json> <out.xlsx>
"""
import json, sys, re
from collections import Counter
from wordfreq import zipf_frequency as zf
from kiwipiepy import Kiwi
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

exp = json.load(open(sys.argv[1], encoding="utf-8"))
secs = json.load(open(sys.argv[2], encoding="utf-8"))
out = sys.argv[3]

def zipf(w):
    c = [w]
    if w.endswith("하다"):
        c += [w[:-2], w[:-2] + "한", w[:-2] + "하게"]
    elif w.endswith("다"):
        c += [w[:-1] + "는", w[:-1] + "게"]
    return round(max(zf(x, "ko") for x in c), 2)

# 사업계획서 등장 문서 수 (형태소 원형 기준)
kiwi = Kiwi()
docfreq = Counter()
for d in secs:
    txt = "\n".join(d[s] for s in ("ITEM", "MOTIV", "DOM", "FOR"))
    if not txt.strip():
        continue
    seen = set()
    for t in kiwi.tokenize(re.sub(r"\n", " ", txt)):
        tag = t.tag.split("-")[0]
        seen.add(t.form + "다" if tag in ("VV", "VA") else t.form)
        if tag == "XR":
            seen.add(t.form + "하다")
    docfreq.update(seen)
NDOC = sum(1 for d in secs if any(d[s] for s in ("ITEM", "MOTIV", "DOM", "FOR")))

def corpus_df(w):
    base = w[:-2] if w.endswith("하다") else w
    return max(docfreq.get(w, 0), docfreq.get(base, 0))

A = "Arial"
hf = Font(name=A, bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="1F3864")
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
base = Font(name=A, size=10)
dimfill = {"Logos": "DDEBF7", "Pathos": "FCE4D6", "Ethos": "E2EFDA"}
wb = Workbook()

ws = wb.active; ws.title = "README"
lines = [
 ("연역 후보: 25개 테마 씨앗 → 표준국어대사전 확장 (Rodale 단계의 한국어판)", True), ("", False),
 ("자료: 국립국어원 표준국어대사전 전체 내려받기 JSON (2026-10-12판), 표제어 362,810개·뜻 509,143개.", False),
 ("경로 A 비슷한말: 씨앗에서 테마에 맞는 뜻만 고른 뒤(씨앗 시트), 그 뜻의 '비슷한말'을 1단계, 1단계 단어의 비슷한말을 2단계로 확장.", False),
 ("경로 B 뜻풀이: 뜻풀이 구절의 끝(머리말)이 씨앗 표현인 표제어. 예) 환희 '매우 기뻐함. 또는 큰 기쁨' → '기쁨'으로 끝나므로 후보. 4어절 이하 구절만, 전문 분야 뜻·참조 표제어(→) 제외.", False),
 ("공통 필터: 일반어, 명사·동사·형용사·부사·어근, 2음절 이상.", False),
 ("왜 두 경로인가: 표준국어대사전은 뜻 가운데 4%에만 비슷한말이 달려 있어 경로 A만으로는 대부분의 씨앗에서 후보가 나오지 않음(기쁨·두려움·효율 등 0개).", False), ("", False),
 ("열 설명", True),
 ("현대 사용 빈도(zipf): wordfreq 한국어 빈도(위키백과·자막·웹 기반). 3 이상이면 흔히 쓰는 말, 0이면 현대 텍스트에서 거의 안 쓰는 옛말·희귀 한자어. 동음이의어 빈도가 섞일 수 있음(예: 시구).", False),
 (f"사업계획서 등장 문서 수: DIPS 사업계획서 {NDOC}건 중 그 단어(형태소 원형)가 한 번이라도 나온 문서 수. 0이면 점수 계산에 영향이 없음.", False),
 ("평정자1~3: Logos/Pathos/Ethos/제외. 판정은 세 명이 같은 차원일 때만 그 차원.", False), ("", False),
 ("권장: 현대 빈도 2.5 이상이거나 사업계획서에 1건 이상 나온 단어부터 평정(필터로 거르기). 나머지는 원 논문처럼 전체를 평정할지, 기준을 정해 제외할지 결정 후 방법론에 기록.", False),
 ("다음 단계: 우리말샘(관련어·비슷한말이 더 풍부) 확장 결과를 같은 형식으로 추가 → 귀납 후보와 합쳐 평정.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = Font(name=A, bold=b, size=12 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

ws = wb.create_sheet("연역후보")
cols = [("No", 6), ("차원", 9), ("테마", 17), ("단어", 14), ("경로", 11), ("단계", 6), ("경유", 10), ("뜻풀이", 55),
        ("현대 빈도(zipf)", 10), ("사업계획서 등장 문서 수", 11), ("평정자1", 10), ("평정자2", 10), ("평정자3", 10), ("판정", 10), ("비고", 16)]
for j, (h, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hf; c.fill = hfill; c.border = B
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(j)].width = w
dv = DataValidation(type="list", formula1='"Logos,Pathos,Ethos,제외"', allow_blank=True); ws.add_data_validation(dv)
order = {"Logos": 0, "Pathos": 1, "Ethos": 2}; porder = {"씨앗": 0, "A 비슷한말": 1, "B 뜻풀이": 2}
rows = exp["rows"]
for r in rows:
    r["z"] = zipf(r["word"]); r["df"] = corpus_df(r["word"])
rows.sort(key=lambda r: (order[r["dim"]], r["theme"], porder[r["path"]], -r["df"], -r["z"]))
inp = PatternFill("solid", fgColor="FFFF00"); grey = Font(name=A, size=10, color="808080")
for i, r in enumerate(rows, 2):
    vals = [i - 1, r["dim"], r["theme"], r["word"], r["path"], r["level"], r["via"], r["defn"], r["z"], r["df"]]
    rare = r["z"] < 2.5 and r["df"] == 0
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = grey if rare and j in (4, 8) else base; c.border = B
        c.alignment = Alignment(vertical="top", wrap_text=(j == 8))
    ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=dimfill[r["dim"]])
    for j in (11, 12, 13):
        c = ws.cell(row=i, column=j); c.fill = inp; c.border = B; c.font = base
    f = ws.cell(row=i, column=14, value=f'=IF(COUNTA(K{i}:M{i})<3,"",IF(AND(K{i}=L{i},L{i}=M{i}),K{i},"불일치"))'); f.font = base; f.border = B
    ws.cell(row=i, column=15).border = B
dv.add(f"K2:M{len(rows)+1}"); ws.freeze_panes = "E2"; ws.auto_filter.ref = f"A1:O{len(rows)+1}"

ws = wb.create_sheet("씨앗")
cols = [("차원", 9), ("테마", 18), ("씨앗", 12), ("선택한 뜻 수", 9), ("선택한 뜻풀이", 90)]
for j, (h, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hf; c.fill = hfill; c.border = B
    ws.column_dimensions[get_column_letter(j)].width = w
for i, s in enumerate(exp["seeds"], 2):
    for j, v in enumerate([s["dim"], s["theme"], s["seed"], s["n_senses"], " / ".join(s["senses"]) or "(사전에 없음)"], 1):
        c = ws.cell(row=i, column=j, value=v); c.font = base; c.border = B; c.alignment = Alignment(wrap_text=(j == 5), vertical="top")

ws = wb.create_sheet("테마별_요약")
cols = [("차원", 9), ("테마", 18), ("후보 수", 9), ("씨앗", 7), ("A 비슷한말", 10), ("B 뜻풀이", 9), ("현대 빈도 2.5 이상", 12), ("사업계획서 등장", 12)]
for j, (h, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hf; c.fill = hfill; c.border = B
    c.alignment = Alignment(horizontal="center", wrap_text=True); ws.column_dimensions[get_column_letter(j)].width = w
themes = sorted({(r["dim"], r["theme"]) for r in rows}, key=lambda t: (order[t[0]], t[1]))
L = len(rows) + 1
for i, (dim, th) in enumerate(themes, 2):
    ws.cell(row=i, column=1, value=dim); ws.cell(row=i, column=2, value=th)
    ws.cell(row=i, column=3, value=f'=COUNTIF(연역후보!$C$2:$C${L},B{i})')
    for j, p in ((4, "씨앗"), (5, "A 비슷한말"), (6, "B 뜻풀이")):
        ws.cell(row=i, column=j, value=f'=COUNTIFS(연역후보!$C$2:$C${L},B{i},연역후보!$E$2:$E${L},"{p}")')
    ws.cell(row=i, column=7, value=f'=COUNTIFS(연역후보!$C$2:$C${L},B{i},연역후보!$I$2:$I${L},">=2.5")')
    ws.cell(row=i, column=8, value=f'=COUNTIFS(연역후보!$C$2:$C${L},B{i},연역후보!$J$2:$J${L},">0")')
    for j in range(1, 9):
        ws.cell(row=i, column=j).font = base; ws.cell(row=i, column=j).border = B
n = len(themes) + 2
ws.cell(row=n, column=2, value="합계").font = Font(name=A, bold=True)
for j in range(3, 9):
    col = get_column_letter(j)
    c = ws.cell(row=n, column=j, value=f"=SUM({col}2:{col}{n-1})"); c.font = Font(name=A, bold=True); c.border = B
wb.save(out)
print(out, len(rows), "modern>=2.5:", sum(r["z"] >= 2.5 for r in rows), "in corpus:", sum(r["df"] > 0 for r in rows))
